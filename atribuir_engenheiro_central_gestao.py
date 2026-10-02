"""
Atribui o engenheiro financeiro a propostas já cadastradas em "Em
Aberto" na Central de Gestão (gestão Banco B), usando duas colunas do
Excel exportado da Plataforma (mesmo arquivo usado por
cadastrar_central_gestao.py): "Inspetor" (nome) pra BUSCAR o engenheiro, e
"CPF Inspetor" (CPF) pra CONFIRMAR com certeza qual é o certo. Nome
sozinho nunca decide - só serve pra trazer candidatos.

Adaptado do robô irmão que já faz essa mesma atribuição pro Banco A
(robo-cadastro-banco-a) - mesma mecânica (busca por nome, decide por CPF,
prévia antes de confirmar, reconferência depois de salvar).

Seletores conferidos um a um contra o HTML real do Banco B
(mapear_central_gestao.py, 09/09/2026):

    td.num_proposta                      OK (traz o número limpo)
    input.checkbox-proposta              OK (value = id interno da linha)
    button[onclick="toggleBotoes(id)"]   OK
    #atriengfinanceModal<id>             OK
    #id_pagengenheiro-<id>               OK
    form do modal -> POST /banco_b/engenheirofinanceiro/<id>   OK

O único que NÃO batia era a caixa de busca: no Banco A ela tem
name="search", no Banco B não tem atributo name nenhum (ver
CAMPO_BUSCA).

Por que nome pra buscar e CPF pra decidir (já validado no fluxo do
Banco A): a busca da Central de Gestão (/buscar-pagengenheiros) não indexa CPF,
só nome/e-mail - buscar direto pelo CPF sempre voltou vazio, mesmo pro
CPF certo; buscar pelo nome completo achou o engenheiro na hora. Nome
nunca decide sozinho: entre os candidatos trazidos pela busca por nome,
só atribui se exatamente 1 tiver CPF IDÊNTICO (comparado só pelos
dígitos, sem pontuação) ao CPF do Excel.

SEGURANÇA (isso mexe em pagamento de engenheiro - nunca chuta):
- A busca por nome só TRAZ candidatos - só atribui quando exatamente 1
  desses candidatos tiver CPF IDÊNTICO ao da coluna "CPF Inspetor" do
  Excel. Zero ou mais de um com CPF batendo = pula e marca pra revisão
  manual.
- Modo "prévia" sempre roda primeiro: mostra proposta -> engenheiro
  encontrado (ou motivo de ter pulado) pra CADA linha do Excel, sem
  clicar em nada ainda, e só segue pro envio de verdade se você digitar
  CONFIRMAR.
- Depois de cada Salvar, reconfere a própria tela (busca a proposta de
  novo e olha a coluna Eng) antes de marcar como sucesso - nunca confia
  só no clique.
- Antes de tentar atribuir, também confere se a proposta JÁ está com o
  engenheiro certo - cobre reexecução depois de Ctrl+C, queda de
  conexão etc. sem reenviar Salvar pra quem já estava certo (mesma
  função usada na reconferência de depois, ver verificar_atribuicao).

PARALELISMO: a etapa [4/4] (atribuir + salvar, a que mais demora - cada
Salvar recarrega a página inteira) roda em várias guias ao mesmo tempo,
cada uma com seu próprio navegador Playwright - mesma estratégia
validada no robô irmão do Banco A. Login/resolução de engenheiro/prévia/
confirmação continuam sequenciais, numa única guia; só depois de
confirmado é que as guias paralelas abrem (reusando a sessão já salva).
Ver comum.paralelismo_configurado pra ajustar quantas guias sem editar
o código (padrão: 6 - validado em produção; antes era 3, conservador,
porque isso mexe em pagamento de verdade a cada Salvar).

Antes de atribuir de verdade, confere quem já saiu de "Em Aberto" (lista
inteira lida da tela antes da prévia) e marca como `FORA` - não é erro,
é o estado normal de quem já foi finalizada numa rodada anterior (a
janela do export cobre vários dias, então isso é esperado, não exceção).
Quem não localiza a linha por outro motivo (duplicada, escondida por
filtro, divergência) usa comum.diagnosticar_linha_nao_achada pra separar
isso de um erro de verdade - mesma classificação usada no robô irmão do
Banco A.

Login automático: com CENTRAL_USUARIO/CENTRAL_SENHA configurados, tenta logar
sozinho quando a sessão salva expira (mesmo mecanismo dos outros
scripts, ver comum.py).

Uso: python atribuir_engenheiro_central_gestao.py [caminho\\para\\arquivo.xlsx]
Se não passar o caminho, usa o Excel mais recente em data/exports/.

⚠️ Ainda não rodou contra o sistema real do Banco B - a mecânica vem
do fluxo já validado do Banco A (não é chute), mas o fluxo completo pro
Banco B precisa do primeiro teste real (com $env:HEADLESS="false" e
um Excel pequeno) pra validar antes de confiar nele com volume.
"""

import os
import re
import sys
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime

from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeout

from comum import (
    abrir_pagina,
    termo_busca_listjs,
    Tee,
    modo_automatico,
    diagnosticar_linha_nao_achada,
    pausar_para_usuario,
    pedir_confirmacao,
    headless_configurado,
    achar_coluna_por_cabecalho,
    achar_excel_mais_recente,
    espera_carregamento_gestao,
    bloquear_recursos_visuais,
    abrir_gestao,
    credenciais_configuradas,
    login_automatico,
    paralelismo_configurado,
    url_sem_cache,
)

LOGIN_URL = "https://central-gestao.example.com/login"
# ⚠️ ainda não confirmado contra o sistema real - assume o mesmo padrão
# de URL usado pro Banco A (/gestao-banco-a), trocando o nome do cliente.
GESTAO_BANCO_B_URL = "https://central-gestao.example.com/gestao-banco-b"
# se a URL acima não responder, o robô procura no menu do sistema um
# link cujo href contenha esta palavra (ver comum.abrir_gestao).
PALAVRA_CHAVE_CLIENTE = "banco_b"
# URL de gestão efetivamente em uso nesta execução. Começa igual à
# constante acima e é substituída se a descoberta pelo menu achar outra -
# os recarregamentos no meio do fluxo (um por proposta, já que cada
# Salvar recarrega a página) precisam voltar pra MESMA tela, não pra
# URL que já se sabe que não funciona.
_url_gestao_em_uso = GESTAO_BANCO_B_URL
BUSCAR_ENGENHEIROS_URL = "https://central-gestao.example.com/buscar-pagengenheiros"
PASTA_SCRIPT = os.path.dirname(os.path.abspath(__file__))
EXPORTS_DIR = os.path.join(PASTA_SCRIPT, "data", "exports")
# mesmo arquivo de sessão dos outros scripts da Central de Gestão.
SESSION_FILE = os.path.join(PASTA_SCRIPT, "sessao_central_gestao.json")
LOG_DIR = os.path.join(PASTA_SCRIPT, "logs")

NOME_COLUNA_PROPOSTA = "proposta"
# confirmado nos cabeçalhos reais do Excel do Banco A (mesmo relatório da
# Plataforma, só client diferente): existem DUAS colunas parecidas -
# "CPF Inspetor" (o CPF) e "Inspetor" (o nome) - por isso a busca por
# cabeçalho aqui é EXATA (exato=True), não por substring: "CPF
# Inspetor" também contém a palavra "Inspetor", então uma busca por
# substring bateria na coluna errada (ou nas duas, dependendo da ordem).
NOME_COLUNA_CPF_INSPETOR = "CPF Inspetor"
NOME_COLUNA_NOME_INSPETOR = "Inspetor"

# Caixa "Procurar..." da tela Consultar. Confirmada no HTML real do
# Banco B (mapear_central_gestao.py, 09/09/2026):
#   <input class="form-control form-control-sm shadow-none search"
#          type="search" placeholder="Procurar..." aria-label="search">
# Repare que NÃO existe name="search" - o seletor herdado do Banco A
# (input.search[name='search']) não casaria com nada, e o robô ficaria
# 5 minutos esperando um elemento inexistente antes de desistir.
CAMPO_BUSCA = "input.search[type='search']"

TAMANHO_MINIMO_BUSCA = 3
TIMEOUT_LONGO = 45000
# a tela de Em Aberto (Consultar) já demorou muito além do razoável pra
# ficar pronta em execuções reais no fluxo irmão do Banco A - timeout bem
# folgado sempre que navega pra ela ou espera algo dela aparecer, em vez
# de um valor curto que trava à toa num dia mais lento.
TIMEOUT_PAGINA_EM_ABERTO = 300000
# a página "parece" carregada (a caixa de busca já aparece) bem antes de
# estar pronta de verdade - o JS/jQuery que faz a busca, os modais etc.
# funcionarem ainda está inicializando por trás. Confirmado no fluxo do
# Banco A: mexer cedo demais trava. Pausa fixa, além do timeout de espera
# por elemento acima (esse já passa antes de a página estar realmente
# pronta). Isso roda 1x por proposta (cada Salvar recarrega a página),
# então o valor pesa bastante no tempo total - ajustável sem editar o
# código, ver comum.espera_carregamento_gestao.
ESPERA_CARREGAMENTO_GESTAO = espera_carregamento_gestao()

# quantas guias atribuem engenheiro em paralelo (etapa [4/4], a que mais
# demora - cada Salvar recarrega a página inteira). Cada guia abre seu
# próprio navegador Playwright (nunca compartilha um só entre threads -
# é a forma segura de paralelizar com a API síncrona). Ver
# comum.paralelismo_configurado pra ajustar sem editar aqui.
PARALELISMO = paralelismo_configurado("PARALELISMO_ATRIBUICAO")

HEADLESS = headless_configurado()

# protege os prints de mais de uma guia escrevendo ao mesmo tempo no
# terminal/log (etapa [4/4], em paralelo) - sem isso, linhas de guias
# diferentes podem sair misturadas no meio umas das outras.
_TRAVA_LOG = threading.Lock()


def _log(msg):
    with _TRAVA_LOG:
        print(msg)


def novo_contexto_pagina(browser, com_sessao):
    context = browser.new_context(
        accept_downloads=True,
        viewport={"width": 1920, "height": 1080} if HEADLESS else None,
        no_viewport=None if HEADLESS else True,
        storage_state=SESSION_FILE if com_sessao else None,
    )
    if HEADLESS:
        bloquear_recursos_visuais(context)
    return context, context.new_page()


def _aguardar_pagina_pronta(page):
    """Pausa fixa depois da caixa de busca já estar visível - a página
    "parece" pronta antes de estar (ver ESPERA_CARREGAMENTO_GESTAO).
    Usa _log (protegido por lock) porque também é chamada de dentro de
    guias paralelas (etapa [4/4])."""
    _log(f"      Aguardando {ESPERA_CARREGAMENTO_GESTAO // 1000}s pra página terminar de carregar de verdade...")
    page.wait_for_timeout(ESPERA_CARREGAMENTO_GESTAO)


def _recarregar_pagina_com_seguranca(page):
    """Sai e volta pra Consultar do zero - usado depois de um erro no
    meio do processamento de uma proposta, pra garantir que a próxima
    começa numa tela limpa (nunca deixa uma exceção daqui escapar, o
    pior caso é a próxima proposta também dar erro)."""
    try:
        page.goto(url_sem_cache(_url_gestao_em_uso), timeout=TIMEOUT_PAGINA_EM_ABERTO, wait_until="domcontentloaded")
        page.wait_for_selector(CAMPO_BUSCA, timeout=TIMEOUT_PAGINA_EM_ABERTO)
        _aguardar_pagina_pronta(page)
    except Exception:
        pass


def ler_propostas_e_engenheiros(caminho_arquivo):
    """Lê (Nº Proposta, Nome do Inspetor, CPF do Inspetor) de todas as
    abas do Excel - pula linhas sem os três valores (nada a atribuir
    ali, ou sem como buscar/confirmar direito)."""
    wb = load_workbook(caminho_arquivo, read_only=True, data_only=True)
    pares = []
    for nome_aba in wb.sheetnames:
        ws = wb[nome_aba]
        coluna_proposta = achar_coluna_por_cabecalho(ws, NOME_COLUNA_PROPOSTA)
        coluna_cpf = achar_coluna_por_cabecalho(ws, NOME_COLUNA_CPF_INSPETOR, exato=True)
        coluna_nome = achar_coluna_por_cabecalho(ws, NOME_COLUNA_NOME_INSPETOR, exato=True)
        if coluna_proposta is None or coluna_cpf is None or coluna_nome is None:
            print(f"      [AVISO] Aba '{nome_aba}' não tem coluna de Proposta, Nome e/ou CPF do Inspetor - pulando.")
            continue
        for linha in ws.iter_rows(min_row=2):
            valor_proposta = linha[coluna_proposta - 1].value
            valor_nome = linha[coluna_nome - 1].value
            valor_cpf = linha[coluna_cpf - 1].value
            num_proposta = str(valor_proposta).strip() if valor_proposta not in (None, "") else ""
            nome_inspetor = str(valor_nome).strip() if valor_nome not in (None, "") else ""
            cpf_inspetor = str(valor_cpf).strip() if valor_cpf not in (None, "") else ""
            if num_proposta and nome_inspetor and cpf_inspetor:
                pares.append({
                    "num_proposta": num_proposta,
                    "nome_inspetor": nome_inspetor,
                    "cpf_inspetor": cpf_inspetor,
                })
    return pares


def _normalizar_num_proposta(texto):
    """Tira espaços e o marcador ⭐ que a tela às vezes mostra ao lado
    do número, pra comparar com o valor do Excel sem falso negativo."""
    texto = texto.replace("⭐️", "").replace("⭐", "")
    return re.sub(r"\s+", "", texto).strip()


def _normalizar_cpf(texto):
    """Fica só com os dígitos - o CPF aparece formatado de jeitos
    diferentes no cadastro de engenheiros da Central de Gestão (confirmado
    no fluxo do Banco A: um com pontos e traço, outro só dígitos), então a
    comparação nunca deve depender de pontuação."""
    return re.sub(r"\D", "", str(texto or ""))


def buscar_engenheiros(page, termo):
    """Chama o mesmo endpoint que a busca da tela usa (GET, sem
    precisar digitar nada visualmente) e devolve a lista de candidatos
    (pode vir vazia)."""
    resposta = page.request.get(BUSCAR_ENGENHEIROS_URL, params={"q": termo})
    if not resposta.ok:
        raise RuntimeError(f"busca de engenheiro devolveu status {resposta.status}")
    return resposta.json()


def buscar_candidatos_por_nome(page, nome):
    """Busca candidatos pelo NOME - a busca da Central de Gestão
    (/buscar-pagengenheiros) indexa nome/e-mail, não CPF (confirmado no
    fluxo do Banco A: buscar direto pelo CPF sempre devolveu vazio, mesmo
    pro CPF certo). O nome só serve pra TRAZER candidatos; qual deles é
    o certo de verdade é decidido depois, comparando CPF (ver
    achar_correspondencia_exata_cpf) - nome nunca é o critério final."""
    termo = nome.strip()
    if len(termo) < TAMANHO_MINIMO_BUSCA:
        return []
    return buscar_engenheiros(page, termo)


def achar_correspondencia_exata_cpf(candidatos, cpf_alvo):
    """Só devolve um candidato se exatamente 1 tiver CPF idêntico
    (comparando só os dígitos) ao CPF procurado - None em qualquer
    outro caso (zero ou mais de um), pra nunca chutar."""
    alvo_norm = _normalizar_cpf(cpf_alvo)
    exatos = [c for c in candidatos if _normalizar_cpf(c.get("cpf", "")) == alvo_norm]
    if len(exatos) == 1:
        return exatos[0]
    return None


JS_PROPOSTAS_EM_ABERTO = """
() => {
    const tabela = document.querySelector('#table');
    if (!tabela) return null;
    const limpa = (el) => (el ? el.textContent : '').replace(/\\s+/g, '').replace(/⭐️|⭐/g, '').trim();
    const out = [];
    tabela.querySelectorAll(':scope > tbody.list > tr td.num_proposta').forEach(td => {
        const v = limpa(td);
        if (v) out.push(v);
    });
    return out;
}
"""


def ler_propostas_em_aberto(page):
    """Devolve o conjunto de nºs de proposta presentes na tela de Em
    Aberto (inclusive linhas escondidas por filtro - textContent). None se
    não conseguiu ler a tabela: aí o chamador não classifica nada como
    FORA de antemão e as guias descobrem uma a uma, como antes."""
    try:
        lista = page.evaluate(JS_PROPOSTAS_EM_ABERTO)
    except Exception:
        return None
    if lista is None:
        return None
    return set(lista)


def resolver_engenheiros(page, pares):
    """Pra cada par (proposta, nome + CPF do inspetor), decide o que
    vai acontecer: engenheiro encontrado (com id) pra atribuir, ou
    pulado com o motivo. Busca por NOME (é o que a Central de Gestão indexa)
    e só confirma pelo CPF (exato, nunca por nome - nome só serve pra
    trazer candidatos). Cacheia por CPF (mesmo inspetor se repete em
    várias propostas) pra não repetir a mesma busca."""
    cache = {}
    decisoes = []
    for par in pares:
        cpf = par["cpf_inspetor"]
        if cpf not in cache:
            digitos = _normalizar_cpf(cpf)
            if len(digitos) != 11:
                cache[cpf] = ("PULAR", None, f"CPF '{cpf}' não tem 11 dígitos - não parece válido")
            else:
                try:
                    candidatos = buscar_candidatos_por_nome(page, par["nome_inspetor"])
                except Exception as e:
                    cache[cpf] = ("PULAR", None, f"erro na busca: {e}")
                else:
                    achado = achar_correspondencia_exata_cpf(candidatos, cpf)
                    if achado is None:
                        qtd = len(candidatos)
                        motivo = "nenhum candidato encontrado buscando pelo nome" if qtd == 0 else \
                            f"{qtd} candidato(s) achado(s) buscando pelo nome, nenhum com CPF EXATAMENTE igual"
                        cache[cpf] = ("PULAR", None, motivo)
                    else:
                        cache[cpf] = ("ATRIBUIR", achado, None)

        acao, achado, motivo = cache[cpf]
        decisoes.append({**par, "acao": acao, "engenheiro": achado, "motivo": motivo})
    return decisoes


# Procura a linha entre TODAS as linhas (não só as visíveis) e, se ela
# estiver escondida por algum filtro da tela, revela só ela antes de
# clicar.
#
# ⚠️ Adaptado por precaução do robô irmão do Banco A, onde isso resolveu um
# problema real: lá a tela esconde com display:none as linhas do tipo
# que não está selecionado num filtro (Financiamento/Consórcio) - depois
# de cada Salvar a página volta pro filtro padrão, e toda proposta do
# outro tipo virava "não achei essa proposta pra clicar". Ainda não
# confirmado se a tela do Banco B tem um filtro equivalente que
# escondesse linhas do mesmo jeito - mas a função é inofensiva se não
# houver nada escondido (só não acha nada a revelar) e evita o mesmo
# problema se houver. Mexer no filtro seria chute (o JS não está no
# HTML da página), então aqui só desfazemos o display:none da linha
# alvo - mudança puramente visual, não toca em dado nenhum.
JS_REVELAR_LINHA = """
(numProposta) => {
    const tabela = document.querySelector('#table');
    if (!tabela) return {qtd: 0};
    const limpa = (el) => (el ? el.textContent : '')
        .replace(/⭐️|⭐/g, '').replace(/\s+/g, '').trim();

    const achados = [];
    for (const tr of tabela.querySelectorAll(':scope > tbody.list > tr')) {
        const celula = tr.querySelector('td.num_proposta');
        if (celula && limpa(celula) === numProposta) achados.push(tr);
    }
    if (achados.length !== 1) return {qtd: achados.length};

    const tr = achados[0];
    const chk = tr.querySelector('input.checkbox-proposta');
    if (!chk) return {qtd: 0};
    if (tr.offsetParent === null) tr.style.display = '';
    return {qtd: 1, id: chk.value};
}
"""


def localizar_linha(page, num_proposta):
    """Usa a busca (List.js) da tela Consultar pra filtrar até sobrar
    só a linha da proposta, e devolve o id interno dela (usado nos ids
    de modal/campo) - ou None se não sobrar EXATAMENTE 1 linha com o
    número EXATO (evita pegar a proposta errada por coincidência de
    substring numa busca parcial).

    Preenche o campo com fill() (rápido, atômico) em vez de digitar
    caractere por caractere - confirmado no fluxo do Banco A que digitar
    trava nessa tela (o List.js reindexa centenas de linhas a cada
    tecla, o que trava a interação no meio). fill() não dispara o
    evento 'keyup' que o List.js escuta pra filtrar, então dispara isso
    à mão logo depois."""
    campo_busca = page.locator(CAMPO_BUSCA)
    campo_busca.click(timeout=TIMEOUT_PAGINA_EM_ABERTO)
    # nunca digitar o nº cru: com ponto/hífen a busca da tela não acha
    # nada (ver comum.termo_busca_listjs). A comparação logo abaixo
    # continua sendo pelo nº EXATO.
    campo_busca.fill(termo_busca_listjs(num_proposta), timeout=TIMEOUT_PAGINA_EM_ABERTO)
    campo_busca.dispatch_event("keyup")
    page.wait_for_timeout(800)  # dá tempo do List.js terminar de filtrar as linhas

    linhas_visiveis = page.locator("tbody.list > tr:visible")
    candidatos = []
    for i in range(linhas_visiveis.count()):
        linha = linhas_visiveis.nth(i)
        texto_coluna = _normalizar_num_proposta(linha.locator("td.num_proposta").inner_text())
        if texto_coluna == num_proposta:
            candidatos.append(linha)

    if len(candidatos) == 1:
        return candidatos[0].locator("input.checkbox-proposta").get_attribute("value")

    # nenhuma linha VISÍVEL bateu: pode ser uma proposta escondida por
    # algum filtro da tela (ver JS_REVELAR_LINHA). Procura entre todas e
    # revela.
    if len(candidatos) == 0:
        achado = page.evaluate(JS_REVELAR_LINHA, num_proposta)
        if achado.get("qtd") == 1:
            page.wait_for_timeout(200)
            return achado["id"]
    return None


def abrir_modal_atribuir_engenheiro(page, linha_id):
    """Clica na seta (mostra os botões escondidos) e depois no botão
    'Atribuir Engenheiro' dessa linha específica."""
    page.locator(f'button[onclick="toggleBotoes({linha_id})"]').click(timeout=TIMEOUT_LONGO)
    botao_atribuir = page.locator(f'button[data-bs-target="#atriengfinanceModal{linha_id}"]')
    botao_atribuir.wait_for(state="visible", timeout=TIMEOUT_LONGO)
    botao_atribuir.click(timeout=TIMEOUT_LONGO)
    modal = page.locator(f"#atriengfinanceModal{linha_id}")
    modal.wait_for(state="visible", timeout=TIMEOUT_LONGO)
    return modal


def atribuir_e_salvar(page, linha_id, engenheiro_id):
    """Com o modal já aberto: preenche o ID do engenheiro (mesmo campo
    que o botão 'Selecionar' da tela preencheria) e clica no Salvar
    DESSE modal - a página tem vários botões 'Salvar', por isso o
    clique é escopado dentro do modal certo.

    Espera só por 'domcontentloaded' (não o padrão 'load') - confirmado
    no fluxo do Banco A que o envio funcionou (o engenheiro ficou salvo de
    verdade) mas o expect_navigation esperando 'load' estourou o
    timeout mesmo assim nessa página pesada; 'domcontentloaded' é o
    sinal que realmente chegou a tempo."""
    page.locator(f"#id_pagengenheiro-{linha_id}").fill(str(engenheiro_id))
    modal = page.locator(f"#atriengfinanceModal{linha_id}")
    botao_salvar = modal.locator("form button[type='submit']")
    with page.expect_navigation(timeout=TIMEOUT_PAGINA_EM_ABERTO, wait_until="domcontentloaded"):
        botao_salvar.click(timeout=TIMEOUT_LONGO)


def _sem_espacos_repetidos(texto):
    return " ".join(str(texto or "").split()).lower()


def verificar_atribuicao(page, num_proposta, nome_engenheiro_esperado):
    """Busca a proposta de novo do zero (a página pode ter recarregado
    depois do Salvar) e confere se o nome do engenheiro esperado
    aparece na linha - nunca confia só em 'o clique não deu erro'.

    Compara ignorando espaços repetidos. Casos reais (29/09 e
    02/10/2026): engenheiros cadastrados na Central com espaço duplo no
    nome ("Cícero Dantas  BA", "Neto -  CarIús CE") voltavam como
    "salva mas não confirmada", porque a tela (HTML) mostra um espaço só.
    Nos dois casos a atribuição tinha dado certo (a Etapa 5 finalizou com
    esse engenheiro). A mesma função decide "já está atribuído" antes de
    salvar, então o falso negativo também fazia salvar de novo à toa."""
    linha_id = localizar_linha(page, num_proposta)
    if linha_id is None:
        return False
    linha = page.locator(f'input.checkbox-proposta[value="{linha_id}"]').locator("xpath=ancestor::tr")
    texto_linha = linha.inner_text()
    return _sem_espacos_repetidos(nome_engenheiro_esperado) in _sem_espacos_repetidos(texto_linha)


def atribuir_lote_paralelo(indice, minha_lista):
    """Atribui uma fatia das decisões (só as com acao == 'ATRIBUIR') numa
    guia própria, em paralelo com as outras. Cada guia abre seu próprio
    navegador Playwright do zero (não compartilha `browser` com as
    outras) - é a forma segura de paralelizar com a API síncrona.

    A sessão salva (SESSION_FILE) e a URL de gestão (_url_gestao_em_uso)
    já foram confirmadas na fase sequencial (login + descoberta da URL +
    resolver engenheiros), antes de qualquer guia abrir - por isso cada
    guia pode reusá-las direto, sem precisar negociar login ou
    redescobrir a URL sozinha.

    Devolve a lista de resultados (mesmo formato usado no Excel final:
    dict com num_proposta, cpf_inspetor, engenheiro, resultado, detalhe)
    dessa fatia - nunca lança exceção pra fora (erro vira resultado
    ERRO/VERIFICAR na lista, nunca derruba a execução inteira)."""
    prefixo = f"[Guia {indice + 1}]"
    resultados_meus = []

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=HEADLESS,
            args=[] if HEADLESS else ["--start-maximized"],
        )
        _contexto, page = novo_contexto_pagina(browser, com_sessao=True)
        try:
            try:
                page.goto(url_sem_cache(_url_gestao_em_uso), timeout=TIMEOUT_PAGINA_EM_ABERTO, wait_until="domcontentloaded")
                page.wait_for_selector(CAMPO_BUSCA, timeout=TIMEOUT_PAGINA_EM_ABERTO)
                _aguardar_pagina_pronta(page)
            except Exception as e:
                _log(f"{prefixo} [ERRO] Não consegui abrir Em Aberto nesta guia ({e}) - "
                     f"marcando as {len(minha_lista)} proposta(s) dela pra revisão manual.")
                return [
                    {**d, "resultado": "ERRO", "detalhe": f"guia não conseguiu abrir a tela: {e}"}
                    for d in minha_lista
                ]

            for d in minha_lista:
                num_proposta = d["num_proposta"]
                eng = d["engenheiro"]
                _log(f"{prefixo} [{num_proposta}] Atribuindo {eng['name']} (id {eng['id']})...")

                try:
                    ja_atribuido = verificar_atribuicao(page, num_proposta, eng["name"])
                except Exception:
                    ja_atribuido = False

                if ja_atribuido:
                    _log(f"{prefixo}    Já está atribuído com {eng['name']} - pulando (nada a fazer).")
                    resultados_meus.append({**d, "resultado": "OK", "detalhe": "já estava atribuído antes desta execução"})
                    continue

                try:
                    linha_id = localizar_linha(page, num_proposta)
                except Exception as e:
                    _log(f"{prefixo}    [ERRO] Não consegui nem localizar a proposta na tela: {e}")
                    resultados_meus.append({**d, "resultado": "ERRO", "detalhe": f"erro localizando a proposta: {e}"})
                    _recarregar_pagina_com_seguranca(page)
                    continue

                if linha_id is None:
                    try:
                        diag = diagnosticar_linha_nao_achada(page, num_proposta)
                        codigo, motivo = diag.get("codigo", "?"), diag.get("texto", "?")
                    except Exception as e:
                        codigo, motivo = "erro_diagnostico", f"não consegui nem diagnosticar ({e})"

                    # "não está mais em Em Aberto" é o estado NORMAL de
                    # quem já foi finalizada numa rodada anterior, não um
                    # erro - importa desde que a janela do export passou a
                    # cobrir 30 dias (exportar_status_banco_b.DIAS_JANELA_AUTOMATICA):
                    # a maioria das propostas do Excel já pode estar
                    # pronta, e chamar cada uma delas de "erro" esconde os
                    # poucos erros de verdade no meio de muitas linhas.
                    if codigo == "fora_da_lista":
                        _log(f"{prefixo}    [OK] Nada a fazer: {motivo}")
                        resultados_meus.append({**d, "resultado": "FORA", "detalhe": motivo})
                    elif codigo in ("duplicada", "visivel_divergente", "escondida"):
                        _log(f"{prefixo}    [PULADO] {motivo}")
                        resultados_meus.append({**d, "resultado": "PULADO", "detalhe": motivo})
                    else:
                        _log(f"{prefixo}    [ERRO] Não achei essa proposta pra clicar: {motivo}")
                        resultados_meus.append({**d, "resultado": "ERRO",
                                                "detalhe": f"proposta não localizada na tela: {motivo}"})
                    continue

                try:
                    abrir_modal_atribuir_engenheiro(page, linha_id)
                    atribuir_e_salvar(page, linha_id, eng["id"])
                except Exception as e:
                    _log(f"{prefixo}    [ERRO] Falha tentando salvar: {e}")
                    resultados_meus.append({**d, "resultado": "ERRO", "detalhe": f"erro ao salvar: {e}"})
                    _recarregar_pagina_com_seguranca(page)
                    continue

                try:
                    # SEMPRE volta pra lista por uma URL sem cache, em vez
                    # de reaproveitar a página que o envio deixou - o
                    # navegador pode servir a cópia ANTERIOR ao
                    # salvamento (confirmado no fluxo do Banco A, onde isso
                    # fazia atribuições bem-sucedidas serem lidas como
                    # "não confirmadas" - ver comum.url_sem_cache).
                    page.goto(url_sem_cache(_url_gestao_em_uso),
                              timeout=TIMEOUT_PAGINA_EM_ABERTO, wait_until="domcontentloaded")
                    page.wait_for_selector(CAMPO_BUSCA, timeout=TIMEOUT_PAGINA_EM_ABERTO)
                    _aguardar_pagina_pronta(page)
                    confirmado_na_tela = verificar_atribuicao(page, num_proposta, eng["name"])
                except Exception as e:
                    _log(f"{prefixo}    [AVISO] Salvei, mas deu erro técnico reconferindo ({e}) - confira manualmente.")
                    resultados_meus.append({**d, "resultado": "VERIFICAR", "detalhe": f"salvo, mas erro técnico ao reconferir: {e}"})
                    _recarregar_pagina_com_seguranca(page)
                    continue

                if confirmado_na_tela:
                    _log(f"{prefixo}    Confirmado na tela - engenheiro atribuído com sucesso.")
                    resultados_meus.append({**d, "resultado": "OK", "detalhe": "confirmado na tela"})
                else:
                    _log(f"{prefixo}    [AVISO] Salvei, mas não vi o nome na tela depois - confira manualmente.")
                    resultados_meus.append({**d, "resultado": "VERIFICAR", "detalhe": "salvo, mas nome não apareceu na tela depois"})
        finally:
            browser.close()

    return resultados_meus


def _tentar_login_automatico(page):
    """Tenta logar sozinho com as credenciais de ambiente (CENTRAL_USUARIO /
    CENTRAL_SENHA). Só devolve True se realmente SAIU da tela de login -
    nunca confia em "o clique não deu erro". Sem credenciais
    configuradas, nem tenta."""
    usuario, senha = credenciais_configuradas("CENTRAL")
    if not usuario:
        return False
    print("      Sessão expirada - tentando login automático...")
    if not login_automatico(page, usuario, senha):
        print("      [AVISO] Não achei o formulário de login pra preencher sozinho.")
        return False
    try:
        page.wait_for_url(lambda url: "/login" not in url, timeout=30000)
    except PlaywrightTimeout:
        pass
    if "/login" in page.url:
        print("      [AVISO] Preenchi o login mas continuo na tela de login "
              "(credencial errada, ou o sistema pediu verificação extra).")
        return False
    return True


def main():
    os.makedirs(LOG_DIR, exist_ok=True)

    log_path = os.path.join(LOG_DIR, f"atribuir_engenheiro_{datetime.now():%Y%m%d_%H%M%S}.txt")
    log_file = open(log_path, "w", encoding="utf-8")
    stdout_original = sys.stdout
    sys.stdout = Tee(stdout_original, log_file)

    try:
        print("=" * 60)
        print(" ATRIBUIR ENGENHEIRO - CENTRAL DE GESTÃO (BANCO B)")
        print("=" * 60)
        print(f"Log desta execução: {log_path}")

        caminho_arquivo = sys.argv[1] if len(sys.argv) > 1 else achar_excel_mais_recente(EXPORTS_DIR)
        if not caminho_arquivo or not os.path.exists(caminho_arquivo):
            print("\n[ERRO] Não achei nenhum Excel pra usar.")
            print("Exporte primeiro com exportar_status_banco_b.py, ou passe o caminho:")
            print(f"  python {os.path.basename(__file__)} caminho\\para\\arquivo.xlsx")
            return
        print(f"Arquivo lido: {caminho_arquivo}")

        pares = ler_propostas_e_engenheiros(caminho_arquivo)
        if not pares:
            print("\nNenhuma linha com Nº Proposta, Nome e CPF do Inspetor preenchidos - nada a fazer.")
            return
        print(f"      {len(pares)} linha(s) com proposta + nome/CPF do inspetor no Excel.")

        with sync_playwright() as p:
            browser = p.chromium.launch(
                headless=HEADLESS,
                args=[] if HEADLESS else ["--start-maximized"],
            )
            sessao_existente = os.path.exists(SESSION_FILE)
            context, page = novo_contexto_pagina(browser, com_sessao=sessao_existente)

            def salvar_sessao():
                context.storage_state(path=SESSION_FILE)

            decisoes = []
            resultados = []

            try:
                print("\n[1/4] Verificando sessão salva...")
                abrir_pagina(page, LOGIN_URL)
                if sessao_existente and "/login" not in page.url:
                    print("      Sessão válida, painel carregado direto.")
                elif _tentar_login_automatico(page):
                    salvar_sessao()
                    print("      Login automático OK - sessão renovada e salva.")
                elif HEADLESS or modo_automatico():
                    print("\n[ERRO] A sessão salva expirou e não consegui logar sozinho.")
                    print("Pra rodar sem ninguém na frente, configure as credenciais uma vez:")
                    print('  [Environment]::SetEnvironmentVariable("CENTRAL_USUARIO", "voce@exemplo.com", "User")')
                    print('  [Environment]::SetEnvironmentVariable("CENTRAL_SENHA", "suasenha", "User")')
                    print("Ou rode com a janela visível e logue na mão:")
                    print(f'  $env:HEADLESS="false"; python {os.path.basename(__file__)}')
                    return
                else:
                    pausar_para_usuario(
                        "A sessão salva não está mais válida (ou é a primeira vez).",
                        "1. Faça o login manualmente na janela do Chrome (e-mail e senha).",
                        "2. Espere a página carregar depois do login (a tela do Dashboard",
                        "   aparecer de vez) ANTES de voltar aqui e apertar ENTER.",
                    )
                    if "/login" in page.url:
                        try:
                            page.wait_for_url(lambda url: "/login" not in url, timeout=15000)
                        except PlaywrightTimeout:
                            pass
                    if "/login" in page.url:
                        print("\n[ERRO] Ainda estou vendo a tela de login.")
                        print("Encerrando esta execução.")
                        return
                    salvar_sessao()
                    print("      Sessão salva para as próximas execuções.")

                print("[2/4] Abrindo Vistorias > Banco B > Em Aberto (Consultar)...")
                global _url_gestao_em_uso
                url_gestao = abrir_gestao(
                    page, GESTAO_BANCO_B_URL, PALAVRA_CHAVE_CLIENTE, timeout=TIMEOUT_PAGINA_EM_ABERTO
                )
                if url_gestao is None:
                    print("\n[ERRO] Não consegui chegar na tela de gestão do Banco B.")
                    print("Encerrando esta execução sem mexer em nada.")
                    return
                # os recarregamentos (um por proposta) precisam voltar pra
                # esta mesma tela, não pra URL padrão que pode não servir.
                _url_gestao_em_uso = url_gestao
                if "/login" in page.url:
                    print("\n[ERRO] Cheguei na tela de gestão mas voltei pra tela de login.")
                    print("Encerrando esta execução.")
                    return
                page.wait_for_selector(CAMPO_BUSCA, timeout=TIMEOUT_PAGINA_EM_ABERTO)
                _aguardar_pagina_pronta(page)

                print("[3/4] Buscando cada Inspetor pelo nome e conferindo se o CPF bate com")
                print("      exatamente 1 engenheiro cadastrado na Central de Gestão (sem clicar em nada ainda)...")
                decisoes = resolver_engenheiros(page, pares)

                # O Excel pode cobrir 30 dias (ver DIAS_JANELA_AUTOMATICA em
                # exportar_status_banco_b.py); parte dele já pode ter
                # sido finalizada e não estar mais em Em Aberto. Sem esta
                # checagem a prévia contaria como "receberia engenheiro"
                # proposta que nem está mais na tela, e cada guia gastaria
                # uma busca inteira só pra descobrir isso na hora.
                em_aberto = ler_propostas_em_aberto(page)
                if em_aberto is not None:
                    for d in decisoes:
                        if d["acao"] == "ATRIBUIR" and d["num_proposta"] not in em_aberto:
                            d["acao"] = "FORA"
                            d["motivo"] = "não está mais em Em Aberto (já finalizada ou movida de fila)"
                    print(f"      {len(em_aberto)} proposta(s) em Em Aberto na tela agora.")
                else:
                    print("      [AVISO] Não consegui ler a lista de Em Aberto de antemão - "
                          "as guias vão descobrir proposta por proposta.")

                total_atribuir = sum(1 for d in decisoes if d["acao"] == "ATRIBUIR")
                total_fora = sum(1 for d in decisoes if d["acao"] == "FORA")
                total_pular = len(decisoes) - total_atribuir - total_fora

                print("\n" + "-" * 60)
                print(" PRÉVIA - nada foi salvo ainda")
                print("-" * 60)
                for d in decisoes:
                    if d["acao"] == "ATRIBUIR":
                        eng = d["engenheiro"]
                        print(f"  [ATRIBUIR] Proposta {d['num_proposta']}: CPF {d['cpf_inspetor']} "
                              f"-> {eng['name']} (id {eng['id']}, {eng.get('uf', '?')}/{eng.get('cidade', '?')})")
                    elif d["acao"] == "FORA":
                        continue  # pode ser muita linha; só a contagem abaixo
                    else:
                        print(f"  [PULAR]    Proposta {d['num_proposta']}: CPF {d['cpf_inspetor']} -> {d['motivo']}")
                print("-" * 60)
                print(f"  {total_atribuir} proposta(s) receberiam engenheiro, {total_pular} seriam puladas,")
                print(f"  {total_fora} já não estão em Em Aberto (nada a fazer).")

                if total_atribuir == 0:
                    print("\nNada pra atribuir de verdade - encerrando sem mexer em nada.")
                    return

                confirmado = pedir_confirmacao(
                    f"{total_atribuir} proposta(s) vão ter o engenheiro atribuído DE VERDADE agora.",
                    f"{total_pular} serão puladas (veja os motivos acima); "
                    f"{total_fora} já não estão em Em Aberto.",
                )
                if not confirmado:
                    print("\nCancelado pelo usuário - nada foi salvo.")
                    return

            except PlaywrightTimeout as e:
                print("\n[ERRO] O robô travou esperando um elemento aparecer na tela.")
                print("Copie a mensagem abaixo e me envie, junto com o que estava")
                print("aparecendo na janela do Chrome nesse momento:")
                print(str(e))
                decisoes = []
            except Exception as e:
                print("\n[ERRO INESPERADO]")
                print(str(e))
                decisoes = []
            finally:
                try:
                    salvar_sessao()
                except Exception:
                    pass
                if not HEADLESS and not modo_automatico():
                    print("\n" + "#" * 60)
                    print("#  A AÇÃO É SUA AGORA - O ROBÔ ESTÁ PAUSADO")
                    print("#" * 60)
                    input(">>> Pressione ENTER aqui para fechar o navegador... ")
                browser.close()

        # a partir daqui, cada guia abre seu próprio navegador Playwright
        # (etapa [4/4], em paralelo) - a sessão usada por elas já foi
        # confirmada/salva na fase sequencial acima, e a URL de gestão
        # já foi descoberta (_url_gestao_em_uso).
        atribuir_lista = [d for d in decisoes if d["acao"] == "ATRIBUIR"]
        resultados = [
            {**d, "resultado": "FORA" if d["acao"] == "FORA" else "PULADO", "detalhe": d["motivo"]}
            for d in decisoes if d["acao"] != "ATRIBUIR"
        ]

        if atribuir_lista:
            num_guias = min(PARALELISMO, len(atribuir_lista))
            partes = [atribuir_lista[i::num_guias] for i in range(num_guias)]
            print(f"\n[4/4] Atribuindo e salvando com {num_guias} guia(s) em paralelo "
                  f"({len(atribuir_lista)} proposta(s) no total)...")

            with ThreadPoolExecutor(max_workers=num_guias) as executor:
                futuro_por_parte = {
                    executor.submit(atribuir_lote_paralelo, i, parte): parte
                    for i, parte in enumerate(partes)
                }
                for futuro in as_completed(futuro_por_parte):
                    try:
                        resultados.extend(futuro.result())
                    except Exception as e:
                        parte_perdida = futuro_por_parte[futuro]
                        print(f"\n[ERRO INESPERADO] Uma guia falhou de forma inesperada ({e}) - "
                              f"marcando as {len(parte_perdida)} proposta(s) dela pra revisão manual.")
                        resultados.extend(
                            {**d, "resultado": "ERRO", "detalhe": f"guia falhou de forma inesperada: {e}"}
                            for d in parte_perdida
                        )

            # reordena igual à prévia - cada guia termina num momento
            # diferente, então a ordem de chegada dos resultados não é
            # a mesma ordem em que as propostas apareceram na prévia.
            ordem = {d["num_proposta"]: i for i, d in enumerate(decisoes)}
            resultados.sort(key=lambda r: ordem.get(r["num_proposta"], 0))

        if resultados:
            caminho_resultado = os.path.join(
                EXPORTS_DIR, f"resultado_engenheiros_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
            )
            wb_resultado = Workbook()
            ws_resultado = wb_resultado.active
            ws_resultado.title = "Resultado"
            ws_resultado.append(["Nº Proposta", "CPF Inspetor (Excel)", "Engenheiro Central", "Resultado", "Detalhe"])
            for r in resultados:
                eng = r.get("engenheiro")
                ws_resultado.append([
                    r["num_proposta"],
                    r["cpf_inspetor"],
                    eng["name"] if eng else "",
                    r["resultado"],
                    r["detalhe"],
                ])
            wb_resultado.save(caminho_resultado)

            conta = lambda tipo: sum(1 for r in resultados if r["resultado"] == tipo)
            total_ok, total_fora = conta("OK"), conta("FORA")
            total_verificar, total_erro, total_pulado = conta("VERIFICAR"), conta("ERRO"), conta("PULADO")
            print("\nConcluído!")
            print("-" * 60)
            print(f"  {total_ok:4d} atribuída(s) e confirmada(s)")
            print(f"  {total_fora:4d} já não estão em Em Aberto (nada a fazer - já finalizadas)")
            if total_verificar:
                print(f"  {total_verificar:4d} salva(s) mas não confirmada(s) - confira na tela")
            if total_pulado:
                print(f"  {total_pulado:4d} pulada(s) - precisam de ação manual (veja os motivos no Excel)")
            if total_erro:
                print(f"  {total_erro:4d} com erro técnico de verdade")
            resolvidas = total_ok + total_fora
            print(f"\n  Nada pendente de ação: {resolvidas} de {len(resultados)}")
            print(f"  Resultado detalhado: {caminho_resultado}")
            print("-" * 60)

    finally:
        sys.stdout = stdout_original
        log_file.close()


if __name__ == "__main__":
    main()
