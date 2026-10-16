"""Etapa 2 (cadastrar_central_gestao.py): não reenviar proposta já
cadastrada, sem estragar o Excel que as Etapas 3 e 4 leem depois.
Também cobre o Excel de uma proposta só do preparar_teste_reenvio.py."""
import os

import pytest
from openpyxl import load_workbook

import cadastrar_central_gestao as cadastrar
import preparar_teste_reenvio

EXPORT = {
    "Inspeções": [
        ["ID", "Nº Proposta", "Status"],
        ["INS1", "10000001", "Laudo Aceito"],
        ["INS2", 10000002, "Laudo Aceito"],      # número em vez de texto no Excel
        ["INS3", "10000003", "Laudo Indeferido"],
        ["INS4", None, "Laudo Aceito"],
    ],
    "Informações": [["Usuário", "teste"]],       # aba sem coluna de proposta
}


@pytest.fixture
def pastas(tmp_path, monkeypatch):
    monkeypatch.setattr(cadastrar, "PROPOSTAS_ENVIADAS_FILE", str(tmp_path / "propostas_ja_enviadas.txt"))
    monkeypatch.setattr(cadastrar, "ENVIOS_DIR", str(tmp_path / "envios"))
    monkeypatch.setattr(preparar_teste_reenvio, "ENVIOS_DIR", str(tmp_path / "envios"))
    return tmp_path


def linhas(caminho, aba="Inspeções"):
    return [list(l) for l in load_workbook(caminho)[aba].iter_rows(min_row=2, values_only=True)]


def test_log_de_enviadas_acumula(pastas):
    """O log local só cresce: uma execução nunca apaga o que a anterior gravou."""
    assert cadastrar.ler_propostas_ja_enviadas() == set()
    cadastrar.salvar_propostas_enviadas({"10000002", "10000001"})
    cadastrar.salvar_propostas_enviadas({"10000003"})
    cadastrar.salvar_propostas_enviadas(set())
    assert cadastrar.ler_propostas_ja_enviadas() == {"10000001", "10000002", "10000003"}


def test_extrair_numeros_proposta(planilha):
    assert cadastrar.extrair_numeros_proposta(planilha(EXPORT)) == {"10000001", "10000002", "10000003"}
    assert cadastrar.extrair_numeros_proposta("export.xls") == set()


def test_filtro_gera_copia_e_nao_toca_no_original(planilha, pastas):
    """No robô irmão do Banco A o filtro salvava por cima do export, e as
    Etapas 3 e 4 só enxergavam as propostas novas da rodada (de 358
    sobravam 36). Aqui o original nunca é tocado."""
    original = planilha(EXPORT)
    copia, total_antes, removidas = cadastrar.filtrar_propostas_ja_enviadas(original, {"10000001", "10000002"})
    assert (total_antes, removidas) == (3, 2)
    assert os.path.dirname(copia) == str(pastas / "envios")
    assert [l[1] for l in linhas(copia)] == ["10000003", None]
    assert len(linhas(original)) == 4


@pytest.mark.parametrize("enviadas, arquivo", [(set(), "export.xlsx"), ({"10000001"}, "export.xls")])
def test_sem_filtro_possivel(enviadas, arquivo, pastas):
    """Primeira execução (nada enviado ainda) ou arquivo antigo .xls: manda como está."""
    assert cadastrar.filtrar_propostas_ja_enviadas(str(pastas / arquivo), enviadas) is None


def test_excel_sem_coluna_de_proposta_nao_e_filtrado(planilha, pastas):
    caminho = planilha({"Aba": [["ID", "Status"], ["INS1", "Laudo Aceito"]]})
    assert cadastrar.filtrar_propostas_ja_enviadas(caminho, {"10000001"}) is None


def test_excel_de_uma_proposta_so(planilha, pastas):
    """Pra testar reenvio: cópia com só a linha da proposta pedida."""
    original = planilha(EXPORT)
    copia = preparar_teste_reenvio.montar_excel_de_uma_proposta(original, "10000003")
    assert copia == str(pastas / "envios" / "teste_10000003.xlsx")
    assert linhas(copia) == [["INS3", "10000003", "Laudo Indeferido"]]
    assert len(linhas(original)) == 4


def test_proposta_inexistente_nao_gera_arquivo(planilha, pastas):
    assert preparar_teste_reenvio.montar_excel_de_uma_proposta(planilha(EXPORT), "99999999") is None
    assert not (pastas / "envios").exists()


def test_avisa_abas_que_a_tela_ignora(planilha):
    """O JS da tela Cadastrar só converte a PRIMEIRA aba: as outras
    sumiriam sem aviso nenhum."""
    assert cadastrar.abas_que_nao_serao_enviadas(planilha(EXPORT)) == ["Informações"]
    assert cadastrar.abas_que_nao_serao_enviadas(planilha({"Inspeções": EXPORT["Inspeções"]})) == []
    assert cadastrar.abas_que_nao_serao_enviadas("export.xls") == []
