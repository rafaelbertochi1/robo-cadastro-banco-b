"""Etapa 1 (exportar_status_banco_b.py): que propostas entram no Excel,
o que é tirado dele e qual período é pedido à Plataforma."""
from datetime import datetime

import pytest
from openpyxl import load_workbook

import exportar_status_banco_b as exportar

CABECALHO = ["Identificador", "Nº Proposta", "Data Status", "Status", "Data Vistoria"]


# --- filtros e Canceladas lidas do Excel ---------------------------------------

def test_filtros_do_relatorio():
    """Todos os tipos menos AVM; Indeferido e Devolvido entram (a vistoria
    foi feita e o engenheiro tem de ser pago)."""
    assert "AVM" not in exportar.TIPOS_INSPECAO
    assert set(exportar.STATUS_INSPECAO) >= {"Laudo Aceito", "Laudo Aceito com Ressalvas",
                                             "Laudo Indeferido", "Laudo Devolvido", "Cancelada"}


def test_canceladas_vem_do_excel(planilha):
    """O Excel é a fonte confiável (a tela só desenha o que cabe na altura
    visível). "Data Status" não pode ser confundida com "Status"."""
    caminho = planilha({"Inspeções": [
        CABECALHO,
        ["INS1", "10000001", "02/09/2026", "Laudo Aceito c/ Ressalva", "01/09/2026"],
        ["INS2", "10000002", "02/09/2026", "Cancelada", "01/09/2026"],
        ["INS3", "10000003", "02/09/2026", "CANCELADA", None],
        [None, "10000004", "02/09/2026", "Cancelada", None],      # sem Identificador: ignorada
    ]})
    assert exportar.ler_canceladas_do_excel(caminho) == ({"INS2", "INS3"}, {"INS2"})


def test_excel_sem_as_colunas_cai_no_modo_antigo(planilha):
    caminho = planilha({"Inspeções": [["ID", "Situação"], ["INS1", "Cancelada"]]})
    assert exportar.ler_canceladas_do_excel(caminho) is None
    assert exportar.ler_canceladas_do_excel("export.xls") is None


# --- limpeza do Excel ---------------------------------------------------------

def test_remove_canceladas_sem_vistoria(planilha):
    caminho = planilha({"Inspeções": [
        CABECALHO,
        ["INS1", "10000001", "02/09/2026", "Laudo Aceito", "01/09/2026"],
        ["INS2", "10000002", "02/09/2026", "Cancelada", None],
        ["INS3", "10000003", "02/09/2026", "Laudo Aceito", "02/09/2026"],
    ]})
    assert exportar.remover_propostas_sem_vistoria(caminho, {"INS2"}) is True
    ws = load_workbook(caminho)["Inspeções"]
    assert [linha[0] for linha in ws.iter_rows(min_row=2, values_only=True)] == ["INS1", "INS3"]


def test_contagem_que_nao_bate_nao_e_confiavel(planilha):
    """Se a 1ª coluna não for o mesmo Identificador da tela, a remoção não
    acha a linha: o Excel não pode seguir sozinho pras próximas etapas."""
    caminho = planilha({"Inspeções": [CABECALHO, ["INS1", "10000001", "02/09/2026", "Laudo Aceito", None]]})
    assert exportar.remover_propostas_sem_vistoria(caminho, {"INS9"}) is False


def test_nada_a_remover_mantem_o_excel():
    assert exportar.remover_propostas_sem_vistoria("qualquer.xlsx", set()) is True
    assert exportar.remover_propostas_sem_vistoria("export.xls", {"INS1"}) is False


def test_remove_aba_informacoes(planilha):
    """A aba de metadados do relatório atrapalha o cadastro na Central."""
    caminho = planilha({"Inspeções": [CABECALHO], "Informações": [["Usuário", "teste"]]})
    exportar.remover_aba_informacoes(caminho)
    assert load_workbook(caminho).sheetnames == ["Inspeções"]


def test_nao_remove_a_unica_aba(planilha):
    caminho = planilha({"Informações": [["Usuário", "teste"]]})
    exportar.remover_aba_informacoes(caminho)
    assert load_workbook(caminho).sheetnames == ["Informações"]


# --- período pedido à Plataforma ----------------------------------------------

@pytest.mark.parametrize("valor, esperado", [
    (None, datetime(2026, 9, 1)),           # padrão: o trabalho começou em setembro/2026
    ("15/09/2026", datetime(2026, 9, 15)),
    ("", None),                             # vazio desliga o piso
    ("2026-09-15", None),                   # formato errado: ignora o piso
])
def test_piso_configurado(monkeypatch, valor, esperado):
    if valor is not None:
        monkeypatch.setenv("DATA_MINIMA", valor)
    assert exportar.piso_configurado() == esperado


@pytest.mark.parametrize("pedida, usada", [
    ("15/08/2026", "01/09/2026"),   # antes do piso: sobe pro piso
    ("10/09/2026", "10/09/2026"),
    ("data errada", "data errada"), # quem valida o formato é a tela da Plataforma
])
def test_aplicar_data_minima(pedida, usada):
    assert exportar.aplicar_data_minima(pedida) == usada


def test_sem_piso_a_data_pedida_vale(monkeypatch):
    monkeypatch.setenv("DATA_MINIMA", "")
    assert exportar.aplicar_data_minima("15/07/2026") == "15/07/2026"


class AgoraFixo(datetime):
    @classmethod
    def now(cls, tz=None):
        return cls(2026, 10, 16, 9, 0)


def test_janela_automatica_termina_hoje(monkeypatch):
    """30 dias terminando hoje (16/10) começam em 17/09."""
    monkeypatch.setattr(exportar, "datetime", AgoraFixo)
    assert exportar.data_inicial_sugerida(exportar.DIAS_JANELA_AUTOMATICA) == "17/09/2026"


def test_janela_automatica_respeita_o_piso(monkeypatch):
    monkeypatch.setattr(exportar, "datetime", AgoraFixo)
    monkeypatch.setenv("DATA_MINIMA", "01/10/2026")
    assert exportar.data_inicial_sugerida(30) == "01/10/2026"
